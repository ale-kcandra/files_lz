import docx # библиотека для работы с вордовскими доками
from collections import Counter # импортируем Counter для подсчета частоты слов и букв
import re # для работы с текстом 
import pandas as pd # подключаем библиотеку pandas для работы с таблицами и данными 
import matplotlib.pyplot as plt # для графика

doc = docx.Document("lion.docx") # открытие дока
num_par = len(doc.paragraphs) # количество абзацев
word_list = [] # список слов без повторений 
a = [] 
rus_let = re.compile(r"[а-яё]", re.IGNORECASE) # регулярное выражение, игнорирующее заглавные буквы
rus_text = [] # список для всего текста
rus_word = re.compile(r"\b[а-яё]+\b", re.IGNORECASE) # регулярное выражение
for paragraph in doc.paragraphs: # для каждого абзаца из абзацев
    text = paragraph.text.lower() # весь текст в нижний регистр
    words = rus_word.findall(text) # находим русские слова и добавляем в список
    letters = rus_let.findall(text) # тоже самое, но с буквами
    a.extend(letters) # добавляем буквы в общий список
    for word in words: # цикл
        rus_text.append(word)
        if word not in word_list: # проверка на наличие слов в списке 
            word_list.append(word) # добавления слова, если его нет в списке 
rus_text_len = len(rus_text) # количество слов в списке
word_counts = Counter(rus_text) # частота встречаемости слов
# создаем словарь, где ключи — это названия колонок для таблицы, а значения — списки с данными
dict_word = {
    "слово": list(word_counts.keys()), # берем все слова из словаря частот и превращаем их в список
    "частота встречи,раз": list(word_counts.values()), # берем все числа и делаем из них список
    "частота встречи в %": [round((count / rus_text_len) * 100, 2) for count in word_counts.values()] # для каждого числа считаем процент от общего количества слов и округляем до 2 знаков
}
df_word = pd.DataFrame(dict_word) # превращаем обычный словарь в таблицу
doc_word = docx.Document() # создаем новый пустой вордовский документ, в который запишем результат
doc_word.add_heading("встречаемость слов в тексте", 0) # добавляем в документ заголовок 

table = doc_word.add_table(rows=1, cols=3) # таблица с одной строкой и тремя колонками 
hdr_cells = table.rows[0].cells # cells - используем для доступа к конкретным ячейкам строки, чтобы заполнить их содержимым
hdr_cells[0].text = "слово" 
hdr_cells[1].text = "частота встречи,раз"
hdr_cells[2].text = "частота встречи в %"

for i in range(len(df_word)):
    row_cells = table.add_row().cells # новые строки в таблицу                                        
    row_cells[0].text = df_word.loc[i, "слово"] # loc - выбор строки с индексом [i]
    row_cells[1].text = str(df_word.loc[i, "частота встречи,раз"])
    row_cells[2].text = str(df_word.loc[i, "частота встречи в %"])
doc_word.save("встречаемость_слов.docx") # сохранение дока 

letters_counts = Counter(a) # частота встречаемости букв
df_letter = pd.DataFrame(letters_counts.items(), columns=["буквы", "количество"])
df_letter["количество"] = pd.to_numeric(df_letter["количество"]) # преобразования значений в столбцах в числовой тип данных
plt.figure(figsize=(10, 6)) # ширина - 10 дюймов, высота - 6
plt.bar(df_letter["буквы"], df_letter["количество"], color="skyblue") # столбчатая диаграмма
plt.xlabel("буквы") # ось X - буквы
plt.ylabel("количество") # ось Y - количество
plt.title("встречаемость букв в тексте") # заголовок для приличия
plt.grid(True) # отображение сетки на графике
plt.show() # вывод графика